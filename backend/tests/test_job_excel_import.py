import io
import uuid
import pytest
import openpyxl
from httpx import AsyncClient, ASGITransport
from app.main import app

@pytest.mark.asyncio
async def test_job_excel_template_and_deduplicated_import():
    uid = uuid.uuid4().hex[:6].upper()
    transport = ASGITransport(app=app)
    headers = {"X-Company-Code": "DEMOLOGISTICS"}

    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Login as Admin
        login = await client.post(
            "/api/v1/auth/login",
            json={"email": "admin@demo.com", "password": "PantherTMS@2026!"},
            headers=headers,
        )
        assert login.status_code == 200
        token = login.json()["access_token"]
        auth_headers = {"Authorization": f"Bearer {token}", "X-Company-Code": "DEMOLOGISTICS"}

        # 2. Test Download Excel Template
        tpl_res = await client.get("/api/v1/transport/jobs/excel-template", headers=auth_headers)
        assert tpl_res.status_code == 200
        assert "application/vnd.openxmlformats" in tpl_res.headers.get("content-type", "")
        tpl_wb = openpyxl.load_workbook(io.BytesIO(tpl_res.content))
        assert tpl_wb.sheetnames == ["Job_Orders_Import"]
        ws = tpl_wb["Job_Orders_Import"]
        assert ws.cell(row=1, column=1).value is not None
        # Assert ZERO demo data: row 2 must be completely empty!
        for col in range(1, 12):
            assert ws.cell(row=2, column=col).value is None

        # 3. Create an in-memory Excel workbook with 3 rows:
        # Row 2: Valid Job A
        # Row 3: Exact duplicate of Job A (Level 1 in-Excel duplicate)
        # Row 4: Valid Job B
        wb = openpyxl.Workbook()
        ws_test = wb.active
        ws_test.title = "Job_Orders_Import"

        # Exactly matches the real form inputs
        headers_list = [
            "Job Number",
            "Billing Client",
            "Origin Location",
            "Destination Location",
            "Date of Job Creation",
            "Scheduled Dispatch Date",
            "Consigner",
            "Consignee",
            "Cargo Description",
            "Estimated Weight",
            "Total Packages",
        ]
        ws_test.append(headers_list)

        job_a_consigner = f"Consigner A {uid}"
        job_a_consignee = f"Consignee A {uid}"
        job_b_consigner = f"Consigner B {uid}"
        job_b_consignee = f"Consignee B {uid}"

        # Row 2: Job A
        ws_test.append([
            "",  # Auto job number
            f"Client Alpha {uid}",
            f"Mumbai {uid}",
            f"Delhi {uid}",
            "2026-09-27",
            "2026-09-28",
            job_a_consigner,
            job_a_consignee,
            "Auto Components",
            "18.5",
            "120",
        ])

        # Row 3: Duplicate of Job A (same client, parties, route, date, cargo)
        ws_test.append([
            "",
            f"Client Alpha {uid}",
            f"Mumbai {uid}",
            f"Delhi {uid}",
            "2026-09-27",
            "2026-09-28",
            job_a_consigner,
            job_a_consignee,
            "Auto Components",
            "18.5",
            "120",
        ])

        # Row 4: Job B (different route and parties)
        ws_test.append([
            f"JOB-CUSTOM-{uid}-1",
            f"Client Beta {uid}",
            f"Pune {uid}",
            f"Bangalore {uid}",
            "2026-09-27",
            "2026-09-29",
            job_b_consigner,
            job_b_consignee,
            "Machinery Spares",
            "24.0",
            "80",
        ])

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        # 4. Upload and Import Excel
        upload_files = {
            "file": ("test_jobs_import.xlsx", excel_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        import_res = await client.post(
            "/api/v1/transport/jobs/import-excel",
            headers=auth_headers,
            files=upload_files,
        )
        assert import_res.status_code == 200, import_res.text
        res_data = import_res.json()

        # Verify Level 1 deduplication:
        # Total rows = 3, imported = 2, skipped duplicates = 1, failed = 0
        assert res_data["total_rows"] == 3
        assert res_data["imported_count"] == 2
        assert res_data["skipped_duplicate_count"] == 1
        assert res_data["failed_count"] == 0
        assert len(res_data["imported_jobs"]) == 2
        assert res_data["skipped_duplicates"][0]["duplicate_type"] == "EXCEL_FILE_DUPLICATE"

        # 5. Upload the EXACT SAME file again to verify Level 2 Database deduplication
        excel_buf.seek(0)
        upload_files_replay = {
            "file": ("test_jobs_import.xlsx", excel_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        replay_res = await client.post(
            "/api/v1/transport/jobs/import-excel",
            headers=auth_headers,
            files=upload_files_replay,
        )
        assert replay_res.status_code == 200, replay_res.text
        replay_data = replay_res.json()

        # Database deduplication must catch both Job A and Job B as existing duplicates!
        assert replay_data["imported_count"] == 0
        assert replay_data["skipped_duplicate_count"] >= 2
        dup_types = [d.get("duplicate_type") for d in replay_data["skipped_duplicates"]]
        assert "DATABASE_DUPLICATE" in dup_types
